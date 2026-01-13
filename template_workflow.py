import re
from docx import Document
from langgraph.graph import StateGraph, END
from neo4j_client import Neo4jClient
from config import settings
from groq import Groq

client = Groq(api_key=settings.GROQ_API_KEY)

# from openai import OpenAI
# client = OpenAI(api_key=settings.OPENAI_API_KEY)


PLACEHOLDER_REGEX = r"<{2,3}\s*[^<>]+?\s*>{2,3}"


def extract_template(state):
    doc = Document(state["template_path"])
    blocks = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            # Extract full placeholder including brackets
            placeholders = re.findall(PLACEHOLDER_REGEX, text)
            blocks.append({"text": text, "placeholders": placeholders})

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells)
            placeholders = re.findall(PLACEHOLDER_REGEX, row_text)
            blocks.append({"text": row_text, "placeholders": placeholders})

    state["blocks"] = blocks
    return state
#--------------------------------------------------------------------------------------------

def fetch_original(state):
    neo = Neo4jClient(
        settings.NEO4J_URI,
        settings.NEO4J_USERNAME,
        settings.NEO4J_PASSWORD
    )
    state["original"] = neo.fetch_document()
    neo.close()
    return state
#--------------------------------------------------------------------------------------------

def resolve_with_llm(state):
    resolved = {}
    original_text = "\n".join([b["text"] for b in state["original"]])

    # Collect all placeholders exactly as in template
    all_placeholders = set()
    for block in state["blocks"]:
        all_placeholders.update(block["placeholders"])

    for placeholder in all_placeholders:
        # All template lines where this placeholder occurs
        related_blocks = [b["text"] for b in state["blocks"] if placeholder in b["text"]]
        block_context = "\n".join(related_blocks)

        # Optimized prompt for legal document extraction
        prompt = f"""
You are an expert at extracting values from legal and financial documents.

TARGET PLACEHOLDER:
{placeholder}

Template lines where it appears (may be repeated):
{block_context}

ORIGINAL DOCUMENT:
{original_text}

INSTRUCTIONS:

1. Find the exact text in the original document that should replace the placeholder.
2. The extracted value must match the format exactly as in the document.
3. Do NOT change the placeholder name or provide any additional text.
4. If the placeholder represents:
   - a NAME → return only the name
   - a DATE → return only the date
   - an AMOUNT → return the numeric and symbol text exactly as it appears
     • Typical amount patterns include currency symbol + digits + commas
     • e.g., "$60,000,000"
   - a NOTE or paragraph → return the full paragraph where this text appears
5. Use the SAME value for all occurrences of this placeholder in the template.
6. Preserve punctuation, formatting, spacing, and capitalization exactly.
7. If the value cannot be found, return NOT_FOUND.
8. Return ONLY the value. Do NOT include examples, comments, or explanation.

Just the extracted value for {placeholder}.
"""

        response = client.chat.completions.create(
        model=settings.GROQ_MODEL,
        messages=[
            {"role": "user", "content": prompt}
        ],
        temperature=0
        )

        value = response.choices[0].message.content.strip()

        if value != "NOT_FOUND":
            resolved[placeholder] = value

    state["resolved"] = resolved
    return state
#--------------------------------------------------------------------------------------------

def fill_template(state):
    

    doc = Document(state["template_path"])
    resolved = state.get("resolved", {})

    def replace_placeholder(paragraph_or_cell, placeholder, value):
        """
        Replace placeholder even if it's split across multiple runs.
        """
        # Combine all runs' text into one string
        full_text = "".join(run.text for run in paragraph_or_cell.runs)
        if placeholder in full_text:
            # Replace all occurrences in full text
            full_text = full_text.replace(placeholder, value)
            # Clear all runs and set full_text in a single run to preserve replacement
            for run in paragraph_or_cell.runs:
                run.text = ""
            paragraph_or_cell.runs[0].text = full_text

    # Replace in paragraphs
    for para in doc.paragraphs:
        for placeholder, value in resolved.items():
            replace_placeholder(para, placeholder, value)

    # Replace in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for placeholder, value in resolved.items():
                        replace_placeholder(para, placeholder, value)

    output_path = "output/FILLED_TEMPLATE.docx"
    doc.save(output_path)
    state["output"] = output_path
    print("*********")
    print("state:===========================", state)
    print("*********")
    return state
#--------------------------------------------------------------------------------------------

def build_workflow():
    g = StateGraph(dict)

    g.add_node("extract_template", extract_template)
    g.add_node("fetch_original", fetch_original)
    g.add_node("resolve_llm", resolve_with_llm)
    g.add_node("fill", fill_template)

    g.set_entry_point("extract_template")
    g.add_edge("extract_template", "fetch_original")
    g.add_edge("fetch_original", "resolve_llm")
    g.add_edge("resolve_llm", "fill")
    g.add_edge("fill", END)

    return g.compile()
#--------------------------------------------------------------------------------------------
