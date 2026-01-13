import os
import uuid
import pdfplumber
from docx import Document
from neo4j_client import Neo4jClient
from config import settings


def extract_original_json(file_path: str) -> dict:
    ext = os.path.splitext(file_path)[1].lower()
    content = []

    if ext in [".doc", ".docx"]:
        doc = Document(file_path)

        for para in doc.paragraphs:
            if para.text.strip():
                content.append({
                    "type": "paragraph",
                    "text": para.text.strip()
                })

        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    content.append({
                        "type": "table_row",
                        "text": " | ".join(cells)
                    })

    elif ext == ".pdf":
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    for line in text.split("\n"):
                        content.append({
                            "type": "paragraph",
                            "text": line.strip()
                        })
    print("content: ",content)
    return {
        "doc_id": str(uuid.uuid4()),
        "content": content
    }
#--------------------------------------------------------------------------------------------

def store_original(file_path: str):
    data = extract_original_json(file_path)

    client = Neo4jClient(
        settings.NEO4J_URI,
        settings.NEO4J_USERNAME,
        settings.NEO4J_PASSWORD
    )
    client.store_document(data["doc_id"], data["content"])
    client.close()

    return data["doc_id"]
#--------------------------------------------------------------------------------------------